# backend/app/rag/ingestion.py
from typing import List, Dict, Any, Optional
from sqlalchemy.orm import Session
import os
import uuid
from datetime import datetime

from ..core.config import settings
from ..models.document import Document
from .chunking import ChunkingService, Chunk
from .embedding import EmbeddingService
from ..utils.logger import logger

class DocumentIngestionService:
    """Service for ingesting documents into the RAG system."""
    
    def __init__(self, db: Session):
        self.db = db
        self.chunking_service = ChunkingService(
            chunk_size=settings.CHUNK_SIZE,
            chunk_overlap=settings.CHUNK_OVERLAP
        )
        self.embedding_service = EmbeddingService(db)
    
    async def ingest_document(
        self,
        document_id: str,
        text: str,
        metadata: Optional[Dict[str, Any]] = None
    ) -> Dict[str, Any]:
        """Ingest a document into the vector store."""
        try:
            # Get document from database
            document = self.db.query(Document).filter(
                Document.id == document_id
            ).first()
            
            if not document:
                raise ValueError(f"Document {document_id} not found")
            
            # Prepare metadata
            doc_metadata = {
                "title": document.title,
                "source_type": document.document_type,
                "university_id": document.university_id,
                "academic_year": document.academic_year,
                "url": document.source_url,
                **(metadata or {})
            }
            
            # Chunk the document
            chunks = self.chunking_service.chunk_text(text, doc_metadata)
            
            if not chunks:
                raise ValueError("No chunks generated from document")
            
            # Generate embeddings for all chunks
            chunk_texts = [chunk.text for chunk in chunks]
            embeddings = await self.embedding_service.generate_embeddings(chunk_texts)
            
            # Store chunks and embeddings
            chunk_ids = []
            for i, (chunk, embedding) in enumerate(zip(chunks, embeddings)):
                chunk_metadata = {
                    **chunk.metadata,
                    "chunk_index": chunk.index,
                    "total_chunks": len(chunks)
                }
                
                chunk_id = self.embedding_service.store_embedding(
                    document_id=document_id,
                    chunk_index=i,
                    chunk_text=chunk.text,
                    embedding=embedding,
                    metadata=chunk_metadata
                )
                chunk_ids.append(chunk_id)
            
            # Update document status
            document.is_indexed = True
            document.chunk_count = len(chunks)
            document.embedding_model = settings.EMBEDDING_MODEL
            document.updated_at = datetime.utcnow()
            
            self.db.commit()
            
            logger.info(f"Successfully ingested document {document_id} with {len(chunks)} chunks")
            
            return {
                "document_id": document_id,
                "chunk_count": len(chunks),
                "chunk_ids": chunk_ids,
                "status": "success"
            }
            
        except Exception as e:
            logger.error(f"Failed to ingest document {document_id}: {str(e)}")
            self.db.rollback()
            raise
    
    async def ingest_file(
        self,
        file_path: str,
        title: str,
        document_type: str,
        university_id: Optional[str] = None,
        academic_year: Optional[str] = None,
        source_url: Optional[str] = None
    ) -> Dict[str, Any]:
        """Ingest a file into the RAG system."""
        # Extract text from file
        text = await self._extract_text(file_path)
        
        # Create document record
        document = Document(
            id=str(uuid.uuid4()),
            title=title,
            document_type=document_type,
            university_id=university_id,
            academic_year=academic_year,
            source_url=source_url,
            file_path=file_path,
            file_size=os.path.getsize(file_path),
            mime_type=self._get_mime_type(file_path),
            is_indexed=False,
            verification_status="PENDING"
        )
        
        self.db.add(document)
        self.db.commit()
        self.db.refresh(document)
        
        # Ingest document
        result = await self.ingest_document(
            document_id=document.id,
            text=text,
            metadata={
                "source_url": source_url,
                "file_name": os.path.basename(file_path)
            }
        )
        
        return {
            "document_id": document.id,
            **result
        }
    
    async def _extract_text(self, file_path: str) -> str:
        """Extract text from a file."""
        ext = os.path.splitext(file_path)[1].lower()
        
        if ext == ".pdf":
            return await self._extract_pdf_text(file_path)
        elif ext == ".docx":
            return await self._extract_docx_text(file_path)
        elif ext == ".txt":
            return await self._extract_txt_text(file_path)
        elif ext == ".csv":
            return await self._extract_csv_text(file_path)
        else:
            raise ValueError(f"Unsupported file type: {ext}")
    
    async def _extract_pdf_text(self, file_path: str) -> str:
        """Extract text from PDF file."""
        try:
            import PyPDF2
            text_parts = []
            with open(file_path, 'rb') as file:
                reader = PyPDF2.PdfReader(file)
                for page_num, page in enumerate(reader.pages):
                    text = page.extract_text()
                    if text:
                        text_parts.append(f"[Page {page_num + 1}]\n{text}")
            return "\n\n".join(text_parts)
        except ImportError:
            raise ImportError("PyPDF2 is required for PDF processing")
    
    async def _extract_docx_text(self, file_path: str) -> str:
        """Extract text from DOCX file."""
        try:
            from docx import Document as DocxDocument
            doc = DocxDocument(file_path)
            text_parts = []
            for para in doc.paragraphs:
                if para.text.strip():
                    text_parts.append(para.text)
            return "\n\n".join(text_parts)
        except ImportError:
            raise ImportError("python-docx is required for DOCX processing")
    
    async def _extract_txt_text(self, file_path: str) -> str:
        """Extract text from TXT file."""
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as file:
            return file.read()
    
    async def _extract_csv_text(self, file_path: str) -> str:
        """Extract text from CSV file."""
        import csv
        rows = []
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as file:
            reader = csv.reader(file)
            for row in reader:
                rows.append(" | ".join(row))
        return "\n".join(rows)
    
    def _get_mime_type(self, file_path: str) -> str:
        """Get MIME type for a file."""
        ext = os.path.splitext(file_path)[1].lower()
        mime_types = {
            ".pdf": "application/pdf",
            ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ".txt": "text/plain",
            ".csv": "text/csv"
        }
        return mime_types.get(ext, "application/octet-stream")